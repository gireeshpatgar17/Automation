"""
Tests for resume text extraction module.
"""

import os
import pytest
from docx import Document
from pypdf import PdfWriter
from resume_matcher.extractor import extract_text, extract_text_from_pdf, extract_text_from_docx


def test_extractor_file_not_found():
    with pytest.raises(FileNotFoundError):
        extract_text_from_pdf("nonexistent.pdf")
    
    with pytest.raises(FileNotFoundError):
        extract_text_from_docx("nonexistent.docx")


def test_extractor_unsupported_format(tmp_path):
    txt_file = tmp_path / "resume.txt"
    txt_file.write_text("Some text")
    with pytest.raises(ValueError):
        extract_text(str(txt_file))


def test_extract_docx(tmp_path):
    docx_file = tmp_path / "resume.docx"
    doc = Document()
    doc.add_paragraph("Jane Doe")
    doc.add_paragraph("Software Engineer")
    doc.save(str(docx_file))

    extracted = extract_text_from_docx(str(docx_file))
    assert "Jane Doe" in extracted
    assert "Software Engineer" in extracted

    # Test via main extract_text function
    extracted_main = extract_text(str(docx_file))
    assert "Jane Doe" in extracted_main
